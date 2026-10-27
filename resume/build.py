"""Build the resume PDF and DOCX from resume.json and publish them to the website.

Usage (from the repository root):
    resume/.venv/bin/python resume/build.py            # build PDF + DOCX, copy to public/resume/
    resume/.venv/bin/python resume/build.py --preview  # also render PNG previews to resume/out/

See resume/README.md for setup.
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import typst
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = HERE / "out"
PUBLIC = ROOT / "public" / "resume"
BASENAME = "Bernard_Abuto_Otieno_Resume"

ACCENT = RGBColor(0x36, 0x55, 0x3F)
MUTED = RGBColor(0x55, 0x54, 0x51)


# --------------------------------------------------------------------- PDF
def build_pdf() -> Path:
    pdf = OUT / f"{BASENAME}.pdf"
    typst.compile(
        str(HERE / "resume.typ"),
        output=str(pdf),
        root=str(HERE),
        font_paths=[str(HERE / "fonts")],
        ignore_system_fonts=True,
    )
    return pdf


def render_previews() -> list[Path]:
    pattern = OUT / "preview-{p}.png"
    typst.compile(
        str(HERE / "resume.typ"),
        output=str(pattern),
        root=str(HERE),
        font_paths=[str(HERE / "fonts")],
        ignore_system_fonts=True,
        format="png",
        ppi=110,
    )
    return sorted(OUT.glob("preview-*.png"))


# --------------------------------------------------------------------- DOCX
def add_hyperlink(paragraph, url: str, text: str, color: RGBColor = ACCENT, size: Pt | None = None):
    """python-docx has no hyperlink API; build the run XML directly."""
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), str(color))
    props.append(c)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size.pt * 2)))
        props.append(sz)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "CFCAC0")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p_pr.append(borders)


def build_docx(data: dict) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(15)
    section.top_margin = section.bottom_margin = Mm(12)
    usable_width = section.page_width - section.left_margin - section.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.05

    h1 = doc.styles["Heading 1"]
    h1.font.name, h1.font.size, h1.font.bold = "Calibri", Pt(11), True
    h1.font.color.rgb = ACCENT
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True

    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.name, bullet_style.font.size = "Calibri", Pt(10)
    bullet_style.paragraph_format.space_after = Pt(1.5)

    def heading(text: str):
        p = doc.add_paragraph(text.upper(), style="Heading 1")
        add_bottom_border(p)

    def entry_line(left_bold: str, left_rest: str, right: str, link: dict | None = None):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        p.add_run(left_bold).bold = True
        p.add_run(left_rest)
        if link:
            p.add_run(" (")
            add_hyperlink(p, link["url"], link["text"])
            p.add_run(")")
        run = p.add_run("\t" + right)
        run.font.color.rgb = MUTED
        return p

    def muted_line(text: str, size: float = 9.5, link: dict | None = None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size, run.font.color.rgb = Pt(size), MUTED
        if link:
            sep = p.add_run("  ·  ")
            sep.font.size, sep.font.color.rgb = Pt(size), MUTED
            add_hyperlink(p, link["url"], link["text"], size=Pt(size))

    # Header
    name = doc.add_paragraph()
    run = name.add_run(data["name"])
    run.bold, run.font.size = True, Pt(20)
    name.paragraph_format.space_after = Pt(1)
    headline = doc.add_paragraph()
    run = headline.add_run(data["headline"])
    run.font.size, run.font.color.rgb = Pt(11.5), ACCENT
    headline.paragraph_format.space_after = Pt(3)
    for line in data["contact"]:
        contact = doc.add_paragraph()
        contact.paragraph_format.space_after = Pt(0)
        for i, item in enumerate(line):
            if i:
                sep = contact.add_run("  ·  ")
                sep.font.size, sep.font.color.rgb = Pt(9.5), MUTED
            if "url" in item and not item["url"].startswith("tel:"):
                add_hyperlink(contact, item["url"], item["text"], size=Pt(9.5))
            else:
                r = contact.add_run(item["text"])
                r.font.size, r.font.color.rgb = Pt(9.5), MUTED

    heading("Summary")
    doc.add_paragraph(data["summary"])

    heading("Experience")
    for job in data["experience"]:
        entry_line(job["title"], f", {job['org']}", job["dates"], job.get("link"))
        for b in job["bullets"]:
            doc.add_paragraph(b, style="List Bullet")

    heading("Selected Projects")
    for p in data["projects"]:
        entry_line(p["name"], f" — {p['focus']}", p["dates"])
        muted_line(p["tech"], link=p["link"])
        for b in p["bullets"]:
            doc.add_paragraph(b, style="List Bullet")

    heading("Education & Training")
    for e in data["education"]:
        entry_line(e["programme"], f", {e['org']}", e["dates"])
        if "detail" in e:
            muted_line(e["detail"])

    heading("Technical Skills")
    for s in data["skills"]:
        p = doc.add_paragraph()
        p.add_run(s["label"] + ": ").bold = True
        p.add_run(s["items"])

    props = doc.core_properties
    props.title = f"{data['name']} — Resume"
    props.author = data["name"]
    props.keywords = "Machine Learning Research Engineer; AI Systems Engineer; NLP; Speech Technology"

    path = OUT / f"{BASENAME}.docx"
    doc.save(path)
    return path


def main() -> None:
    data = json.loads((HERE / "resume.json").read_text(encoding="utf-8"))
    OUT.mkdir(exist_ok=True)
    for old in OUT.glob("preview-*.png"):
        old.unlink()

    pdf = build_pdf()
    docx_path = build_docx(data)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    for path in (pdf, docx_path):
        shutil.copy2(path, PUBLIC / path.name)
        print(f"built   {path.relative_to(ROOT)} -> public/resume/{path.name}")

    if "--preview" in sys.argv:
        for png in render_previews():
            print(f"preview {png.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
